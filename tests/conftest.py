from pathlib import Path

import docx
import pytest

from skillpose.core import prompts


@pytest.fixture(autouse=True)
def _no_limit_from_the_environment(monkeypatch):
    """A MONTHLY_LIMIT_USD of the shell must not change the limit of any test."""
    monkeypatch.delenv("MONTHLY_LIMIT_USD", raising=False)


@pytest.fixture(autouse=True)
def _no_log_file(monkeypatch):
    """A command of a test writes no skillpose.log next to its database; the tests of the log switch it on."""
    monkeypatch.setenv("SKILLPOSE_LOG", "off")


@pytest.fixture(autouse=True)
def _examples_are_the_active_prompts(monkeypatch):
    """Every test starts with the example prompts; one that chooses a set (--prompts, SKILLPOSE_PROMPTS) does not
    change the prompts of the next. A SKILLPOSE_PROMPTS of the shell (private prompts) must not reach a
    test either: the two guards that need it read it once, when their module is imported (before this runs)."""
    monkeypatch.setattr(prompts, "_active", None)
    monkeypatch.delenv("SKILLPOSE_PROMPTS", raising=False)


def make_docx(tmp_path: Path, para: str, cell: str) -> Path:
    doc = docx.Document()
    doc.add_paragraph(para)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = cell
    path = tmp_path / "cv.docx"
    doc.save(path)
    return path


def make_pdf(tmp_path: Path, text: str) -> Path:
    """A minimal one-page PDF with Helvetica text; empty text gives a PDF without a text layer."""
    content = f"BT /F1 10 Tf 40 760 Td ({text}) Tj ET".encode("latin-1") if text else b""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    path = tmp_path / "cv.pdf"
    path.write_bytes(out)
    return path
