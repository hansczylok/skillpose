"""Read CVs and postings from files. Nothing is fetched from the network."""

import csv
import io
import json
import logging
from pathlib import Path

import docx
from pypdf import PdfReader

log = logging.getLogger(__name__)

ALLOWED_CV_TYPES = (".pdf", ".docx", ".txt", ".md")
POSTING_TYPES = (".md", ".txt")
MAX_CV_BYTES = 5 * 1024 * 1024
MIN_CV_CHARS = 200


class CvTextError(ValueError):
    pass


def _plain(path: Path, encoding: str = "utf-8") -> str:
    data = path.read_bytes()
    try:
        return data.decode(encoding)
    except UnicodeDecodeError:
        return data.decode("cp1252")


def _pdf(path: Path) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)


def _docx(path: Path) -> str:
    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def read_cv(path: Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in ALLOWED_CV_TYPES:
        raise CvTextError(f"unsupported file type {suffix or '(none)'} – use {', '.join(ALLOWED_CV_TYPES)}")
    if path.stat().st_size > MAX_CV_BYTES:
        raise CvTextError("file larger than 5 MB")
    text = {".pdf": _pdf, ".docx": _docx}.get(suffix, _plain)(path)
    if len(text.strip()) < MIN_CV_CHARS:
        raise CvTextError("no readable text – scanned PDF? Export the CV as text-based PDF, DOCX or TXT.")
    return text


def _ref(ref, path: Path, position: int) -> str:
    """The given reference, or file name and position (from 1) when the entry has none."""
    return f"{path.name}#{position}" if ref is None or str(ref).strip() == "" else str(ref).strip()


def _json_postings(path: Path) -> list[tuple[str, str]]:
    items = json.loads(path.read_text(encoding="utf-8-sig"))
    return [(_ref(item.get("ref"), path, i), str(item["text"]).strip()) for i, item in enumerate(items, 1)]


def _csv_postings(path: Path) -> list[tuple[str, str]]:
    """A header row with a `text` column and an optional `ref` column. The delimiter (`,` `;` or tab) is guessed
    from the header only: free text in the rows would mislead the guess."""
    text = _plain(path, "utf-8-sig")
    try:
        delimiter = csv.Sniffer().sniff((text.splitlines() or [""])[0], delimiters=",;\t").delimiter
    except csv.Error:
        delimiter = ","
    reader = csv.reader(io.StringIO(text, newline=None), delimiter=delimiter)  # \r\n inside a text becomes \n
    header = [h.strip().lower() for h in next(reader, [])]
    if "text" not in header:
        raise ValueError("CSV needs a 'text' column")
    col = header.index("text")
    ref_col = header.index("ref") if "ref" in header else None
    postings = []
    for position, row in enumerate(reader, 1):
        if not any(cell.strip() for cell in row):
            continue
        if len(row) > len(header):
            if col != len(header) - 1:
                raise ValueError(f"{path.name} line {reader.line_num}: more fields than columns - put the text "
                                 f"in quotes if it contains {delimiter!r}")
            row = [*row[:col], delimiter.join(row[col:])]  # an unquoted delimiter in the last column is text
        cell = row[col] if col < len(row) else ""
        ref = row[ref_col] if ref_col is not None and ref_col < len(row) else None
        postings.append((_ref(ref, path, position), cell.strip()))
    return postings


def read_postings(paths: list[Path]) -> list[tuple[str, str]]:
    """Return (reference, text) per posting: a directory gives its .md/.txt files, a .json file a list of
    {"text": ..., "ref": ...} objects, a .csv file one posting per row (columns `text` and `ref`), a .md/.txt file
    one posting. A JSON or CSV entry without ref is named after its file and position, e.g. `jobs.csv#3`."""
    postings: list[tuple[str, str]] = []
    for path in map(Path, paths):
        suffix = path.suffix.lower()
        if path.is_dir():
            postings.extend((f.name, _plain(f).strip()) for f in sorted(path.iterdir())
                            if f.suffix.lower() in POSTING_TYPES)
        elif suffix == ".json":
            postings.extend(_json_postings(path))
        elif suffix == ".csv":
            postings.extend(_csv_postings(path))
        elif suffix in POSTING_TYPES:
            postings.append((path.name, _plain(path).strip()))
        else:
            log.warning("skipped %s: postings must be .md, .txt, .json or .csv", path)
    return postings
